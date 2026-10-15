"""Grant splits preserve the own-funded cost and the original bank movement."""
from __future__ import annotations

import datetime as dt
import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from finance.models import (
    Account, Company, FundingSource, InternalTransferMatch,
    LedgerClassificationEvent, Transaction, get_default_company,
)
from finance.services import ledger_v2
from finance.services import reports as rep
from finance.views.analytics import _pnl_calendar


class GrantFixture(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_superuser(
            'grant-regression', 'grant@example.test', 'x',
        )
        cls.company = get_default_company()
        cls.account = Account.objects.create(
            company=cls.company, name='Грантовий рахунок', currency='UAH', is_business=True,
        )
        cls.category = cls.company.categories.create(name='Грантове обладнання', type='expense')
        cls.when = timezone.make_aware(dt.datetime(2026, 6, 12, 12))
        cls.source = FundingSource.objects.create(
            company=cls.company, name='Програма 2', received_amount=Decimal('216000'),
        )

    def txn(self, amount, *, kind='operating_expense', ttype=Transaction.TYPE_EXPENSE,
            account=None, when=None, category=None, base=None, **kwargs):
        """Create local ledger facts without invoking bank or FX providers."""
        return Transaction.objects.create(
            company=self.company, type=ttype, status=Transaction.STATUS_ACTUAL,
            amount=Decimal(amount), amount_base=Decimal(base if base is not None else amount),
            currency=(account or self.account).currency, account=account or self.account,
            date_actual=when or self.when, category=category or self.category,
            ownership_scope='business', is_business=True, economic_kind=kind, **kwargs,
        )

    def source_with_budget(self, amount='1000', **kwargs):
        return FundingSource.objects.create(
            company=self.company, name=kwargs.pop('name', 'Інша програма'),
            received_amount=Decimal(amount), **kwargs,
        )

    def allocate(self, txn, amount, source=None, **kwargs):
        return ledger_v2.allocate_funding(
            source=source or self.source, txn=txn, amount=Decimal(amount), user=self.user, **kwargs,
        )

    def reports(self, **params):
        params = {'period': 'custom', 'date_from': '2026-06-01', 'date_to': '2026-06-30', **params}
        return (rep.pnl(self.company, params), rep.cash_flow(self.company, params))

    def assert_report_cost(self, report, expected):
        expected = Decimal(expected)
        self.assertEqual(report.get('expenses', report.get('cash_out')), expected)
        self.assertEqual(
            sum((Decimal(str(row['total'])) for row in report['expense_by_category']), Decimal('0')),
            expected,
        )
        # The calendar/chart consumes this same public series, not the total.
        self.assertEqual(
            sum((Decimal(str(row['out'])) for row in report['series']), Decimal('0')),
            expected,
        )


class GrantReportingTests(GrantFixture):
    def setUp(self):
        self.receipt = self.txn(
            '216000', ttype=Transaction.TYPE_INCOME, kind='grant_inflow', funding_source=self.source,
        )
        self.equipment = self.txn('225870.85', kind='investment', funding_source=self.source)
        self.allocate(self.equipment, '216000')

    @override_settings(ALLOWED_HOSTS=['fin.twocomms.shop', 'testserver'], SECURE_SSL_REDIRECT=False)
    def test_pnl_and_cash_flow_exports_match_report_split_and_include_switch(self):
        self.client.force_login(self.user)
        for kind in ('pnl', 'cashflow'):
            for include, expense, income in (('', '9870.85', '0'), ('1', '225870.85', '216000')):
                with self.subTest(kind=kind, include=include):
                    response = self.client.get(
                        f'/analytic/report/{kind}/export/',
                        data={'period': 'all', 'include_grants': include}, HTTP_HOST='fin.twocomms.shop',
                    )
                    self.assertEqual(response.status_code, 200)
                    workbook = load_workbook(BytesIO(response.content), data_only=True)
                    try:
                        sheet = workbook.active
                        self.assertEqual(sheet['A3'].value, 'Доходи')
                        self.assertEqual(sheet['A4'].value, 'Витрати')
                        self.assertEqual(Decimal(str(sheet['B3'].value)), Decimal(income))
                        self.assertEqual(Decimal(str(sheet['B4'].value)), Decimal(expense))
                        self.assertEqual(Decimal(str(sheet['B5'].value)), Decimal('-9870.85'))
                        self.assertEqual(Decimal(str(sheet['B6'].value)), Decimal('216000'))
                        rows = list(sheet.iter_rows(values_only=True))
                        category_total = sum(
                            (Decimal(str(row[2])) for row in rows if row[0] == 'Витрата'), Decimal('0'),
                        )
                        self.assertEqual(category_total, Decimal(expense))
                        series_header = next(index for index, row in enumerate(rows) if row[0] == 'Період')
                        self.assertEqual(
                            sum((Decimal(str(row[2])) for row in rows[series_header + 1:]), Decimal('0')),
                            Decimal(expense),
                        )
                    finally:
                        workbook.close()

    @override_settings(ALLOWED_HOSTS=['fin.twocomms.shop', 'testserver'], SECURE_SSL_REDIRECT=False)
    def test_raw_transaction_export_keeps_full_bank_amounts(self):
        self.client.force_login(self.user)
        response = self.client.get(
            '/analytic/report/statement/export/', data={'period': 'all'}, HTTP_HOST='fin.twocomms.shop',
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        try:
            rows = list(workbook.active.iter_rows(values_only=True))
            self.assertEqual(rows[0][:3], ('Дата', 'Тип', 'Сума'))
            self.assertEqual(
                sorted(Decimal(str(row[2])) for row in rows[1:]),
                [Decimal('216000'), Decimal('225870.85')],
            )
        finally:
            workbook.close()

    def test_pnl_series_and_calendar_use_agreement_date_cash_flow_uses_actual_date(self):
        actual = timezone.make_aware(dt.datetime(2026, 10, 5, 12))
        agreement = timezone.make_aware(dt.datetime(2026, 10, 3, 12))
        self.equipment.date_actual = actual
        self.equipment.date_agreement = agreement
        self.equipment.save(update_fields=['date_actual', 'date_agreement'])
        params = {'period': 'custom', 'date_from': '2026-10-01', 'date_to': '2026-10-31'}
        pnl, cash = self.reports(**params)
        for report in (pnl, cash):
            self.assert_report_cost(report, '9870.85')
        self.assertEqual(pnl['series'], [{'label': '2026-10-03', 'in': 0.0, 'out': 9870.85}])
        self.assertEqual(cash['series'], [{'label': '2026-10-05', 'in': 0.0, 'out': 9870.85}])
        calendar = _pnl_calendar(self.company, params, 'custom', agreement.date())
        days = {cell['day']: cell for cell in calendar['cells'] if cell.get('day')}
        self.assertEqual(Decimal(str(days[3]['expense'])), Decimal('9870.85'))
        self.assertEqual(Decimal(str(days[5]['expense'])), Decimal('0'))

    def test_partial_grant_excludes_exactly_allocated_share_everywhere(self):
        pnl, cash = self.reports()
        for report in (pnl, cash):
            with self.subTest(report='pnl' if 'expenses' in report else 'cash'):
                self.assert_report_cost(report, '9870.85')
                self.assertEqual(report['targeted_in'], Decimal('216000'))
                self.assertEqual(report['targeted_out'], Decimal('216000'))
        self.assertEqual(pnl['income'], Decimal('0'))
        self.assertEqual(pnl['profit'], Decimal('-9870.85'))
        self.assertEqual(cash['cash_in'], Decimal('0'))
        self.assertEqual(cash['net'], Decimal('-9870.85'))
        self.equipment.refresh_from_db()
        self.assertEqual(self.equipment.amount, Decimal('225870.85'))
        self.assertEqual(self.equipment.amount_base, Decimal('225870.85'))

    def test_include_switch_restores_whole_receipt_and_expense(self):
        pnl, cash = self.reports(include_grants='1')
        for report in (pnl, cash):
            self.assert_report_cost(report, '225870.85')
            self.assertEqual(sum(Decimal(str(row['in'])) for row in report['series']), Decimal('216000'))
        self.assertEqual(pnl['income'], Decimal('216000'))
        self.assertEqual(pnl['profit'], Decimal('-9870.85'))
        self.assertEqual(cash['cash_in'], Decimal('216000'))
        self.assertEqual(cash['net'], Decimal('-9870.85'))

    def test_calendar_uses_same_partial_grant_cost_and_include_switch(self):
        for include, expense, income in (('', '9870.85', '0'), ('1', '225870.85', '216000')):
            with self.subTest(include=include):
                calendar = _pnl_calendar(
                    self.company, {'include_grants': include}, 'custom', self.when.date(),
                )
                day = next(cell for cell in calendar['cells'] if cell.get('day') == 12)
                self.assertEqual(Decimal(str(day['expense'])), Decimal(expense))
                self.assertEqual(Decimal(str(day['income'])), Decimal(income))
                self.assertEqual(Decimal(str(day['profit'])), Decimal('-9870.85'))

    def test_tag_without_allocation_is_not_treated_as_fully_grant_funded(self):
        self.equipment.funding_allocations.all().delete()
        for report in self.reports():
            self.assert_report_cost(report, '225870.85')
            self.assertEqual(report['targeted_out'], Decimal('0'))

    def test_fully_funded_expense_has_zero_cost_without_hiding_bank_row(self):
        self.equipment.amount = self.equipment.amount_base = Decimal('216000')
        self.equipment.save(update_fields=['amount', 'amount_base'])
        for report in self.reports():
            self.assert_report_cost(report, '0')
            self.assertEqual(report['expense_by_category'], [])
        self.assertTrue(Transaction.objects.filter(pk=self.equipment.pk).exists())

    def test_other_funding_type_is_still_a_cost(self):
        self.equipment.funding_allocations.all().delete()
        other = self.source_with_budget('216000', source_type='other')
        self.allocate(self.equipment, '216000', source=other)
        for report in self.reports():
            self.assert_report_cost(report, '225870.85')
            self.assertEqual(report['targeted_out'], Decimal('0'))

    def test_reserved_allocation_and_source_tag_do_not_replace_spent_evidence(self):
        self.equipment.funding_allocations.all().delete()
        self.allocate(self.equipment, '216000', allocation_type='reserved')
        for report in self.reports():
            self.assert_report_cost(report, '225870.85')
            self.assertEqual(report['targeted_out'], Decimal('0'))

    def test_account_category_and_period_filters_preserve_split(self):
        other_account = Account.objects.create(company=self.company, name='Інший', currency='UAH')
        other_category = self.company.categories.create(name='Інше обладнання', type='expense')
        self.txn('800', account=other_account, category=other_category)
        self.txn('900', when=self.when - dt.timedelta(days=40))
        for params in ({'accounts': str(self.account.pk)}, {'categories': str(self.category.pk)}):
            for report in self.reports(**params):
                self.assert_report_cost(report, '9870.85')
                self.assertEqual(report['targeted_out'], Decimal('216000'))
        for report in self.reports(accounts=str(other_account.pk)):
            self.assert_report_cost(report, '800')
            self.assertEqual(report['targeted_out'], Decimal('0'))

    def test_fx_split_rounds_native_remainder_with_decimal_arithmetic(self):
        self.equipment.delete()
        fx_account = Account.objects.create(company=self.company, name='EUR', currency='EUR')
        # 0.01 EUR of a 0.03 EUR transaction remains own-funded: 1.03 / 3 -> 0.34 UAH.
        expense = self.txn('0.03', base='1.03', account=fx_account)
        self.allocate(expense, '0.02')
        for report in self.reports():
            self.assert_report_cost(report, '0.34')
            self.assertEqual(report['targeted_out'], Decimal('0.69'))

    def test_transfer_principal_is_excluded_and_fee_matches_total_chart_and_calendar(self):
        pension = Account.objects.create(company=self.company, name='Пенсійна', currency='UAH')
        debit = self.txn('10050', account=pension, kind='internal_transfer')
        credit = self.txn('10000', ttype=Transaction.TYPE_INCOME, kind='internal_transfer')
        InternalTransferMatch.objects.create(
            company=self.company, source_transaction=debit, destination_transaction=credit,
            source_account=pension, destination_account=self.account,
            amount=Decimal('10000'), fee_amount=Decimal('50'), status='confirmed',
        )
        for report in self.reports():
            self.assert_report_cost(report, '9920.85')
            self.assertEqual(report['transfer_fees'], Decimal('50'))
        calendar = _pnl_calendar(self.company, {}, 'custom', self.when.date())
        day = next(cell for cell in calendar['cells'] if cell.get('day') == 12)
        self.assertEqual(Decimal(str(day['expense'])), Decimal('9920.85'))
        self.assertEqual(Decimal(str(day['profit'])), Decimal('-9920.85'))


class GrantAllocationTests(GrantFixture):
    def setUp(self):
        self.expense = self.txn('1000')

    def test_cross_source_sum_cannot_exceed_transaction(self):
        second = self.source_with_budget()
        self.allocate(self.expense, '700')
        with self.assertRaises(ValueError):
            self.allocate(self.expense, '300.01', source=second)
        self.assertEqual(self.expense.funding_allocations.count(), 1)
        self.allocate(self.expense, '300', source=second)
        self.assertEqual(self.expense.funding_allocations.count(), 2)

    def test_replacement_is_idempotent_and_revalidates_other_sources(self):
        second = self.source_with_budget()
        allocation = self.allocate(self.expense, '700')
        self.allocate(self.expense, '300', source=second)
        for _ in range(2):
            saved = self.allocate(self.expense, '700', replace=True)
            self.assertEqual(saved.pk, allocation.pk)
        with self.assertRaises(ValueError):
            self.allocate(self.expense, '700.01', replace=True)
        allocation.refresh_from_db()
        self.assertEqual(allocation.amount, Decimal('700'))

    def test_grant_cap_includes_existing_spending_and_reservations(self):
        limited = self.source_with_budget('100')
        self.allocate(self.expense, '60', source=limited)
        self.allocate(self.expense, '20', source=limited, allocation_type='reserved')
        other_expense = self.txn('100')
        with self.assertRaises(ValueError):
            self.allocate(other_expense, '20.01', source=limited)
        self.allocate(other_expense, '20', source=limited)
        self.assertEqual(ledger_v2.funding_summary(limited)['available'], Decimal('0'))
        self.allocate(self.expense, '60', source=limited, replace=True)
        self.allocate(self.expense, '50', source=limited, replace=True)
        self.assertEqual(ledger_v2.funding_summary(limited)['available'], Decimal('10'))

    def test_fx_budget_is_base_currency_but_allocations_are_native_currency(self):
        limited = self.source_with_budget('2000')
        fx_account = Account.objects.create(company=self.company, name='USD grant', currency='USD')
        fx_expense = self.txn('100', base='4000', account=fx_account)
        granted = ledger_v2.sync_grant_allocation(fx_expense, source=limited)
        self.assertEqual(granted, Decimal('50'))
        self.assertEqual(fx_expense.funding_allocations.get().amount, Decimal('50'))
        self.assertEqual(ledger_v2.funding_summary(limited)['spent'], Decimal('2000'))
        self.assertEqual(ledger_v2.funding_summary(limited)['available'], Decimal('0'))
        with self.assertRaises(ValueError):
            self.allocate(fx_expense, '50.01', source=limited, replace=True)
        for report in self.reports(accounts=str(fx_account.pk)):
            self.assert_report_cost(report, '2000')

    def test_blank_amount_uses_transaction_limit_when_budget_is_larger(self):
        self.assertEqual(
            ledger_v2.sync_grant_allocation(self.expense, source=self.source), Decimal('1000'),
        )
        self.assertEqual(self.expense.funding_allocations.get().amount, Decimal('1000'))

    def test_sync_blank_amount_uses_available_budget_and_repeats_without_growth(self):
        limited = self.source_with_budget('600')
        for amount in (None, ''):
            granted = ledger_v2.sync_grant_allocation(self.expense, source=limited, amount=amount)
            self.assertEqual(granted, Decimal('600'))
            self.assertEqual(list(self.expense.funding_allocations.values_list('amount', flat=True)), [Decimal('600')])

    def test_sync_switch_and_clear_remove_stale_spent_allocations_only(self):
        second = self.source_with_budget('500')
        self.allocate(self.expense, '700')
        self.allocate(self.expense, '10', allocation_type='reserved')
        ledger_v2.sync_grant_allocation(self.expense, source=second, amount='400')
        self.assertEqual(
            list(self.expense.funding_allocations.filter(allocation_type='spent').values_list('funding_source_id', 'amount')),
            [(second.pk, Decimal('400'))],
        )
        ledger_v2.sync_grant_allocation(self.expense)
        self.assertFalse(self.expense.funding_allocations.filter(allocation_type='spent').exists())
        self.assertTrue(self.expense.funding_allocations.filter(allocation_type='reserved', amount=10).exists())

    def test_invalid_sync_amount_restores_previous_allocation(self):
        original = self.allocate(self.expense, '700')
        second = self.source_with_budget('500')
        for amount in ('500.01', '-1', 'NaN', 'Infinity'):
            with self.subTest(amount=amount):
                with self.assertRaises(ValueError):
                    ledger_v2.sync_grant_allocation(self.expense, source=second, amount=amount)
                original.refresh_from_db()
                self.assertEqual(original.amount, Decimal('700'))
                self.assertEqual(self.expense.funding_allocations.count(), 1)


@override_settings(ALLOWED_HOSTS=['testserver', 'fin.twocomms.shop'], SECURE_SSL_REDIRECT=False)
class GrantClassificationApiTests(GrantFixture):
    def setUp(self):
        self.expense = self.txn('225870.85', kind='investment')
        self.client.force_login(self.user)

    def classify(self, **payload):
        return self.client.post(
            f'/api/v2/transactions/{self.expense.pk}/classification/',
            data=json.dumps({
                'economic_kind': 'investment', 'ownership_scope': 'business',
                'sync_grant_allocation': True, **payload,
            }), content_type='application/json', HTTP_HOST='fin.twocomms.shop',
        )

    def test_checkbox_automatically_caps_grant_and_preserves_own_cost(self):
        for _ in range(2):
            response = self.classify(funding_source_id=self.source.pk)
            self.assertEqual(response.status_code, 200, response.content)
            self.assertEqual(
                list(self.expense.funding_allocations.values_list('amount', flat=True)), [Decimal('216000')],
            )
        self.expense.refresh_from_db()
        self.assertEqual(self.expense.funding_source_id, self.source.pk)
        for report in self.reports():
            self.assert_report_cost(report, '9870.85')

    def test_checkbox_switch_and_clear_update_reports_without_stale_split(self):
        self.assertEqual(self.classify(funding_source_id=self.source.pk).status_code, 200)
        second = self.source_with_budget('100000')
        response = self.classify(funding_source_id=second.pk, funding_amount='90000')
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(
            list(self.expense.funding_allocations.values_list('funding_source_id', 'amount')),
            [(second.pk, Decimal('90000'))],
        )
        for report in self.reports():
            self.assert_report_cost(report, '135870.85')
        response = self.classify(funding_source_id=None)
        self.assertEqual(response.status_code, 200, response.content)
        self.expense.refresh_from_db()
        self.assertIsNone(self.expense.funding_source_id)
        self.assertFalse(self.expense.funding_allocations.exists())
        for report in self.reports():
            self.assert_report_cost(report, '225870.85')

    def test_failed_allocation_rolls_back_classification_category_and_audit(self):
        self.assertEqual(self.classify(funding_source_id=self.source.pk).status_code, 200)
        self.expense.refresh_from_db()
        original_category = self.expense.category_id
        event_count = LedgerClassificationEvent.objects.count()
        original_allocation = self.expense.funding_allocations.get()
        limited = self.source_with_budget('10')
        response = self.classify(
            funding_source_id=limited.pk, funding_amount='11', economic_kind='operating_expense',
            ownership_scope='personal',
        )
        self.assertEqual(response.status_code, 400, response.content)
        self.expense.refresh_from_db()
        self.expense.ledger_classification.refresh_from_db()
        self.assertEqual(self.expense.economic_kind, 'investment')
        self.assertEqual(self.expense.ownership_scope, 'business')
        self.assertEqual(self.expense.ledger_classification.economic_kind, 'investment')
        self.assertEqual(self.expense.funding_source_id, self.source.pk)
        self.assertEqual(self.expense.category_id, original_category)
        self.assertEqual(LedgerClassificationEvent.objects.count(), event_count)
        original_allocation.refresh_from_db()
        self.assertEqual(original_allocation.amount, Decimal('216000'))

    def test_unknown_or_foreign_source_rejects_without_erasing_existing_split(self):
        self.assertEqual(self.classify(funding_source_id=self.source.pk).status_code, 200)
        foreign_company = Company.objects.create(name='Other tenant', base_currency='UAH')
        foreign = FundingSource.objects.create(
            company=foreign_company, name='Foreign grant', received_amount=Decimal('216000'),
        )
        for source_id in (foreign.pk, foreign.pk + 10000):
            with self.subTest(source_id=source_id):
                response = self.classify(funding_source_id=source_id)
                self.assertEqual(response.status_code, 400, response.content)
                self.expense.refresh_from_db()
                self.assertEqual(self.expense.funding_source_id, self.source.pk)
                self.assertEqual(
                    list(self.expense.funding_allocations.values_list('amount', flat=True)), [Decimal('216000')],
                )
